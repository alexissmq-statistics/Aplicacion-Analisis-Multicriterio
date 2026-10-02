"""
=========================================================================================
NÚCLEO DE CÁLCULO: MÉTODO AHP (SAATY) - MODOS "OBJETIVO" Y "SUBJETIVO"
=========================================================================================
Este módulo contiene ÚNICAMENTE funciones de cálculo puro (reciben datos, devuelven
resultados). No usa input(), print(), plt.show() ni google.colab, para que pueda
importarse limpiamente desde una app Streamlit (o desde cualquier otro front-end).

Incluye dos modos, replicando los dos sistemas del notebook original:

  - MODO "AHP" (objetivo):  las matrices de alternativas se generan automáticamente
    a partir de datos reales de desempeño (costo/beneficio).
  - MODO "AHP2" (subjetivo): las matrices de alternativas se arman con comparación
    pareada subjetiva (igual que la matriz de criterios).

El núcleo matemático (normalización, vector de prioridades, consistencia, síntesis y
sensibilidad) es común a ambos modos.
"""

import io
from datetime import datetime

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

# =======================================================================================
# 1. CONSTANTES MATEMÁTICAS: ÍNDICE DE ALEATORIEDAD (RI) DE SAATY
# =======================================================================================
RI_DICT = {1: 0.00, 2: 0.00, 3: 0.58, 4: 0.90, 5: 1.12, 6: 1.24, 7: 1.32, 8: 1.41, 9: 1.45, 10: 1.49}


# =======================================================================================
# 2. UTILIDADES DE CONSTRUCCIÓN DE MATRICES
# =======================================================================================

def construir_matriz_pareada(n, valores_triangulares):
    """
    Construye una matriz recíproca de Saaty (n x n) a partir de los valores del
    triángulo superior.

    Parámetros
    ----------
    n : int
        Dimensión de la matriz (número de criterios o alternativas).
    valores_triangulares : list[float]
        Lista con los valores en orden (0,1), (0,2), ..., (0,n-1), (1,2), ..., (n-2,n-1).
        Longitud esperada: n*(n-1)/2.

    Retorna
    -------
    np.ndarray : matriz n x n con diagonal 1 y reciprocidad garantizada.
    """
    esperado = n * (n - 1) // 2
    if len(valores_triangulares) != esperado:
        raise ValueError(
            f"Se esperaban {esperado} valores para n={n}, se recibieron {len(valores_triangulares)}."
        )
    matriz = np.ones((n, n))
    k = 0
    for i in range(n):
        for j in range(i + 1, n):
            val = float(valores_triangulares[k])
            if val == 0:
                raise ValueError("Los valores de comparación no pueden ser 0.")
            matriz[i, j] = val
            matriz[j, i] = 1.0 / val
            k += 1
    return matriz


def validar_matriz_cuadrada_y_reciproca(matriz, atol=1e-3):
    """
    Verifica que la matriz sea cuadrada y recíproca (a_ij * a_ji ≈ 1).
    Retorna (es_valida: bool, mensaje: str).
    """
    sh = matriz.shape
    if sh[0] != sh[1]:
        return False, "La matriz no es cuadrada."
    for i in range(sh[0]):
        for j in range(sh[0]):
            if not np.isclose(matriz[i, j], 1.0 / matriz[j, i], atol=atol):
                return False, f"Falla de reciprocidad matemática en la coordenada ({i},{j})."
    return True, "Matriz válida."


# =======================================================================================
# 3. NÚCLEO ALGORÍTMICO: PRIORIDADES Y CONSISTENCIA
# =======================================================================================

def procesar_matriz_prioridades(matriz, nombres_filas, nombres_columnas):
    """
    Implementa los pasos del AHP para obtener los pesos de prioridad:
    Suma de columnas -> Normalización -> Promedio por filas.

    Retorna un diccionario con cada paso intermedio como DataFrame, más el
    vector de pesos final (np.ndarray), para que la interfaz pueda mostrar
    el desarrollo paso a paso si lo desea.
    """
    df_original = pd.DataFrame(matriz, index=nombres_filas, columns=nombres_columnas)

    sumas_columnas = np.sum(matriz, axis=0)
    df_sumas = pd.DataFrame([sumas_columnas], columns=nombres_columnas, index=["SUMA"])

    matriz_normalizada = matriz / sumas_columnas
    df_normalizada = pd.DataFrame(matriz_normalizada, index=nombres_filas, columns=nombres_columnas)

    vector_pesos = np.mean(matriz_normalizada, axis=1)
    df_pesos = pd.DataFrame(vector_pesos, index=nombres_filas, columns=["Peso (W)"])

    return {
        "matriz_original": df_original,
        "sumas_columnas": df_sumas,
        "matriz_normalizada": df_normalizada,
        "vector_pesos_df": df_pesos,
        "vector_pesos": vector_pesos,
    }


def calcular_consistencia(matriz, vector_pesos):
    """
    Calcula Lambda Máximo, Índice de Consistencia (CI) y Razón de Consistencia (CR).
    Retorna (lambda_max, ci, cr, estado, vector_verificacion, lambdas_locales).
    """
    n = matriz.shape[0]
    if n <= 2:
        return n, 0.0, 0.0, "CONSISTENTE (Por propiedad dimensional n<=2)", None, None

    vector_verificacion = np.dot(matriz, vector_pesos)
    lambdas_locales = vector_verificacion / vector_pesos
    lambda_max = np.mean(lambdas_locales)
    ci = (lambda_max - n) / (n - 1)
    ri = RI_DICT.get(n, 1.49)
    cr = ci / ri if ri > 0 else 0.0
    estado = "CONSISTENTE" if cr <= 0.10 else "INCONSISTENTE (Requiere revisión de juicios)"

    return lambda_max, ci, cr, estado, vector_verificacion, lambdas_locales


# =======================================================================================
# 3.5 CARGA DE DATOS DESDE ARCHIVO CSV (SIN google.colab, recibe el contenido ya leído)
# =======================================================================================

def _limpiar_lineas(contenido):
    """Divide el contenido de un CSV en líneas limpias (sin \\r, sin líneas vacías)."""
    return [line.strip() for line in contenido.split("\n") if line.strip()]


def cargar_datos_ahp_csv(contenido):
    """
    Parsea un CSV con el formato del MODO "AHP" (objetivo / costo-beneficio).
    Retorna (criterios, tipos, matriz_crit, alternativas, df_rendimiento).
    """
    lines = _limpiar_lineas(contenido)
    if len(lines) < 4:
        raise ValueError("El archivo CSV es demasiado corto. Se esperan al menos 4 líneas para la configuración inicial.")

    n_crit = int(lines[0].split(",")[1])
    n_alt = int(lines[1].split(",")[1])

    criterios = [c.strip() for c in lines[2].split(",")[1:n_crit + 1]]
    tipos_raw = [t.strip().upper() for t in lines[3].split(",")[1:n_crit + 1]]
    tipos = ["Beneficio" if t == "B" else "Costo" for t in tipos_raw]

    idx_matriz_start = 4
    expected_crit_matrix_end = idx_matriz_start + n_crit
    if len(lines) < expected_crit_matrix_end:
        raise ValueError(
            f"Faltan líneas para la matriz de criterios. Se esperaban {n_crit} filas "
            f"comenzando en la línea {idx_matriz_start + 1}."
        )

    matriz_crit = np.zeros((n_crit, n_crit))
    for i in range(n_crit):
        row_vals = lines[idx_matriz_start + i].split(",")[1:n_crit + 1]
        matriz_crit[i, :] = [float(eval(v)) for v in row_vals]  # eval maneja strings tipo '1/3'

    idx_rend_start = idx_matriz_start + n_crit
    expected_rend_matrix_end = idx_rend_start + n_alt
    if len(lines) < expected_rend_matrix_end:
        actual_alt_rows = len(lines) - idx_rend_start
        raise ValueError(
            f"Se declararon {n_alt} alternativas pero solo se encontraron {actual_alt_rows} "
            f"filas de datos de rendimiento a partir de la línea {idx_rend_start + 1}."
        )

    alternativas = []
    datos_rendimiento = []
    for i in range(n_alt):
        parts = lines[idx_rend_start + i].split(",")
        alternativas.append(parts[0].strip())
        datos_rendimiento.append([float(v) for v in parts[1:n_crit + 1]])

    df_rendimiento = pd.DataFrame(datos_rendimiento, index=alternativas, columns=criterios)
    return criterios, tipos, matriz_crit, alternativas, df_rendimiento


def cargar_datos_ahp2_csv(contenido):
    """
    Parsea un CSV con el formato del MODO "AHP2" (subjetivo / comparación pareada).
    Retorna (criterios, alternativas, matriz_crit, matrices_alternativas).
    """
    lines = _limpiar_lineas(contenido)
    if len(lines) < 4:
        raise ValueError("El archivo CSV es demasiado corto para el formato AHP2.")

    n_crit = int(lines[0].split(",")[1])
    n_alt = int(lines[1].split(",")[1])

    criterios = [c.strip() for c in lines[2].split(",")[1:n_crit + 1]]
    alternativas = [a.strip() for a in lines[3].split(",")[1:n_alt + 1]]

    def extraer_numeros_linea(linea, tamano_esperado):
        partes = [p.strip() for p in linea.split(",") if p.strip()]
        try:
            float(eval(partes[0]))
        except (ValueError, SyntaxError, NameError, ZeroDivisionError):
            partes = partes[1:]  # descartar etiqueta/nombre si no es numérico
        partes = partes[:tamano_esperado]
        return [float(eval(p)) for p in partes]

    try:
        idx_mat_crit_header = next(i for i, l in enumerate(lines) if l.strip() == "Matriz_Criterios")
    except StopIteration:
        raise ValueError("No se encontró la sección 'Matriz_Criterios' en el archivo CSV.")
    idx_mat_crit = idx_mat_crit_header + 1

    if len(lines) < idx_mat_crit + n_crit:
        raise ValueError("Faltan filas para completar la 'Matriz_Criterios'.")

    matriz_crit = np.zeros((n_crit, n_crit))
    for i in range(n_crit):
        matriz_crit[i, :] = extraer_numeros_linea(lines[idx_mat_crit + i], n_crit)

    matrices_alternativas = {}
    idx_actual = idx_mat_crit + n_crit
    while idx_actual < len(lines):
        if lines[idx_actual].startswith("Matriz_Alternativas_"):
            header = lines[idx_actual]
            criterio_nombre = header.replace("Matriz_Alternativas_", "").strip()
            if len(lines) < idx_actual + 1 + n_alt:
                raise ValueError(f"Faltan filas para completar la matriz de alternativas de '{criterio_nombre}'.")
            mat_alt = np.zeros((n_alt, n_alt))
            for i in range(n_alt):
                mat_alt[i, :] = extraer_numeros_linea(lines[idx_actual + 1 + i], n_alt)
            matrices_alternativas[criterio_nombre] = mat_alt
            idx_actual += n_alt + 1
        else:
            idx_actual += 1

    faltantes = [c for c in criterios if c not in matrices_alternativas]
    if faltantes:
        raise ValueError(
            f"No se encontraron matrices de alternativas para los criterios: {faltantes}. "
            "Verifica que cada sección se llame exactamente 'Matriz_Alternativas_<nombre_criterio>'."
        )

    return criterios, alternativas, matriz_crit, matrices_alternativas


# =======================================================================================
# 4. MODO "AHP" (OBJETIVO): MATRICES DE ALTERNATIVAS DESDE DATOS REALES
# =======================================================================================

def generar_matrices_alternativas_objetivas(df_rendimiento, tipos_criterios):
    """
    Construye las matrices de comparación pareada para las alternativas a partir
    de sus ratios reales de desempeño, según orientación Costo o Beneficio.

    df_rendimiento : pd.DataFrame (alternativas x criterios) con valores numéricos reales.
    tipos_criterios : list[str] con "Beneficio" o "Costo" para cada criterio, en el
                      mismo orden que las columnas de df_rendimiento.
    """
    alternativas = list(df_rendimiento.index)
    criterios = list(df_rendimiento.columns)
    n_alt = len(alternativas)

    matrices_alternativas = {}
    for idx_crit, criterio in enumerate(criterios):
        tipo = tipos_criterios[idx_crit]
        matriz_alt = np.ones((n_alt, n_alt))
        for i in range(n_alt):
            for j in range(n_alt):
                val_i = df_rendimiento.iloc[i, idx_crit]
                val_j = df_rendimiento.iloc[j, idx_crit]
                if tipo == "Beneficio":
                    matriz_alt[i, j] = val_i / val_j
                else:
                    matriz_alt[i, j] = val_j / val_i
        matrices_alternativas[criterio] = matriz_alt

    return matrices_alternativas


# =======================================================================================
# 5. SÍNTESIS GLOBAL Y RANKING (COMÚN A AMBOS MODOS)
# =======================================================================================

def calcular_sintesis_y_ranking(pesos_criterios, df_prioridades_locales):
    """
    Multiplica las prioridades locales de las alternativas por los pesos globales
    de los criterios para obtener el ranking final.

    df_prioridades_locales : pd.DataFrame (alternativas x criterios), cada celda es
                              el peso local de esa alternativa bajo ese criterio.
    pesos_criterios : array-like alineado con las columnas de df_prioridades_locales.

    Retorna (df_ranking, mejor_opcion).
    """
    puntaje_global = np.dot(df_prioridades_locales.values, pesos_criterios)

    df_ranking = pd.DataFrame(
        {"Puntaje Global": puntaje_global}, index=df_prioridades_locales.index
    )
    df_ranking = df_ranking.sort_values(by="Puntaje Global", ascending=False)
    df_ranking["Ranking"] = range(1, len(df_ranking) + 1)

    mejor_opcion = df_ranking.index[0]
    return df_ranking, mejor_opcion


# =======================================================================================
# 6. ANÁLISIS DE SENSIBILIDAD (SIN input(), RECIBE LOS CAMBIOS YA DECIDIDOS)
# =======================================================================================

def ejecutar_analisis_sensibilidad(pesos_originales, df_prioridades_locales, cambios_objetivo):
    """
    Redistribuye proporcionalmente el peso remanente entre los criterios no fijados.

    pesos_originales : array-like, pesos actuales de los criterios (en el mismo orden
                        que las columnas de df_prioridades_locales).
    df_prioridades_locales : pd.DataFrame (alternativas x criterios).
    cambios_objetivo : dict {indice_criterio: nuevo_peso}. El resto se redistribuye.

    Retorna un diccionario con:
      - pesos_modificados (np.ndarray)
      - df_comparativo_criterios (DataFrame)
      - df_resultados (DataFrame con Score y Rank originales vs modificados)
      - ranking_estable (bool)
    """
    criterios = list(df_prioridades_locales.columns)
    n_criterios = len(criterios)
    pesos_originales = np.array(pesos_originales, dtype=float)

    if len(cambios_objetivo) >= n_criterios:
        raise ValueError(
            "No se pueden fijar manualmente todos los criterios; al menos uno debe "
            "quedar libre para absorber el residuo matemático."
        )

    suma_pesos_fijos = sum(cambios_objetivo.values())
    for idx, peso in cambios_objetivo.items():
        if peso < 0 or peso > 1:
            raise ValueError("Cada peso debe estar en el intervalo [0, 1].")
    if suma_pesos_fijos >= 1.0:
        raise ValueError(
            f"La suma de los pesos fijos ({suma_pesos_fijos:.4f}) es >= 1.0; "
            "no queda margen para balancear los criterios restantes."
        )

    pesos_modificados = np.zeros(n_criterios)
    for idx, peso in cambios_objetivo.items():
        pesos_modificados[idx] = peso

    indices_restantes = [i for i in range(n_criterios) if i not in cambios_objetivo]
    suma_original_restantes = sum(pesos_originales[i] for i in indices_restantes)
    peso_remanente_nuevo = 1.0 - suma_pesos_fijos

    for i in indices_restantes:
        if suma_original_restantes > 0:
            pesos_modificados[i] = pesos_originales[i] * (peso_remanente_nuevo / suma_original_restantes)
        else:
            pesos_modificados[i] = peso_remanente_nuevo / len(indices_restantes)

    assert np.isclose(np.sum(pesos_modificados), 1.0), "Error interno: la suma de pesos modificados no es 1.0"

    puntaje_original = np.dot(df_prioridades_locales.values, pesos_originales)
    puntaje_modificado = np.dot(df_prioridades_locales.values, pesos_modificados)

    df_comparativo_criterios = pd.DataFrame(
        {"Peso Original": pesos_originales, "Peso Modificado": pesos_modificados}, index=criterios
    )

    df_resultados = pd.DataFrame(
        {"Score Orig.": puntaje_original, "Score Modif.": puntaje_modificado},
        index=df_prioridades_locales.index,
    )
    df_resultados["Rank Orig."] = df_resultados["Score Orig."].rank(ascending=False).astype(int)
    df_resultados["Rank Modif."] = df_resultados["Score Modif."].rank(ascending=False).astype(int)

    ranking_estable = df_resultados["Rank Orig."].equals(df_resultados["Rank Modif."])

    return {
        "pesos_modificados": pesos_modificados,
        "df_comparativo_criterios": df_comparativo_criterios,
        "df_resultados": df_resultados.sort_values(by="Rank Modif."),
        "ranking_estable": ranking_estable,
    }


# =======================================================================================
# 7. GENERACIÓN DE REPORTE DE RESULTADOS (TEXTO PLANO, DESCARGABLE)
# =======================================================================================

def generar_reporte_texto(pesos_criterios, criterios, df_ranking, mejor_opcion, resultado_sens=None):
    """
    Arma un reporte en texto plano con: pesos de criterios, ranking final, mejor
    alternativa y (si se ejecutó) el resultado del análisis de sensibilidad.

    resultado_sens : dict o None. Si se provee, debe tener la misma forma que retorna
                      ejecutar_analisis_sensibilidad().
    """
    lineas = []
    lineas.append("=" * 70)
    lineas.append("REPORTE DE RESULTADOS - SISTEMA AHP")
    lineas.append("=" * 70)

    lineas.append("\n--- PESOS DE LOS CRITERIOS ---")
    for c, w in zip(criterios, pesos_criterios):
        lineas.append(f"  {c}: {w:.4f}")

    lineas.append("\n--- RANKING FINAL DE ALTERNATIVAS ---")
    lineas.append(df_ranking.round(4).to_string())

    lineas.append(f"\n>>> MEJOR ALTERNATIVA: {mejor_opcion}")
    lineas.append(f"    Puntaje global: {df_ranking.iloc[0]['Puntaje Global']:.4f}")

    if resultado_sens is not None:
        lineas.append("\n" + "-" * 70)
        lineas.append("ANÁLISIS DE SENSIBILIDAD")
        lineas.append("-" * 70)
        lineas.append("\nComparativa de pesos (original vs. modificado):")
        lineas.append(resultado_sens["df_comparativo_criterios"].round(4).to_string())
        lineas.append("\nImpacto en el ranking de alternativas:")
        lineas.append(resultado_sens["df_resultados"].to_string())
        if resultado_sens["ranking_estable"]:
            lineas.append("\nConclusión: el ranking se mantuvo ESTABLE ante el cambio de pesos.")
        else:
            lineas.append("\nConclusión: el ranking CAMBIÓ ante el cambio de pesos.")

    lineas.append("\n" + "=" * 70)
    return "\n".join(lineas)


# =======================================================================================
# 8. GENERACIÓN DE REPORTE EN EXCEL (RANKING + MEJOR ALTERNATIVA + GRÁFICOS + SENSIBILIDAD)
# =======================================================================================

_FUENTE = "Arial"
_AZUL_OSCURO = "1E293B"
_AZUL = "4A90E2"
_VERDE = "2ECC71"
_GRIS_CLARO = "F1F5F9"
_BLANCO = "FFFFFF"


def _estilo_encabezado(celda):
    celda.font = Font(name=_FUENTE, bold=True, color=_BLANCO, size=11)
    celda.fill = PatternFill("solid", fgColor=_AZUL_OSCURO)
    celda.alignment = Alignment(horizontal="center", vertical="center")


def _escribir_df(ws, df, fila_inicio, incluir_index=True, nombre_index="Alternativa",
                  resaltar_primera_fila=False):
    """Escribe un DataFrame en una hoja de openpyxl con formato básico. Retorna la
    fila siguiente disponible después de la tabla."""
    col_inicio = 1
    columnas = ([nombre_index] if incluir_index else []) + list(df.columns)

    for j, nombre_col in enumerate(columnas):
        celda = ws.cell(row=fila_inicio, column=col_inicio + j, value=nombre_col)
        _estilo_encabezado(celda)

    borde_fino = Border(*[Side(style="thin", color="CBD5E1")] * 4)

    for i, (idx, fila) in enumerate(df.iterrows()):
        fila_excel = fila_inicio + 1 + i
        valores = ([idx] if incluir_index else []) + list(fila.values)
        for j, val in enumerate(valores):
            celda = ws.cell(row=fila_excel, column=col_inicio + j, value=val)
            celda.font = Font(name=_FUENTE, size=10)
            celda.border = borde_fino
            if isinstance(val, float):
                celda.number_format = "0.0000"
            if resaltar_primera_fila and i == 0:
                celda.fill = PatternFill("solid", fgColor=_VERDE)

    for j, nombre_col in enumerate(columnas):
        ancho = max(12, len(str(nombre_col)) + 4)
        ws.column_dimensions[get_column_letter(col_inicio + j)].width = ancho

    return fila_inicio + 1 + len(df) + 2


def generar_reporte_excel(pesos_criterios, criterios, df_ranking, mejor_opcion,
                           modo, fig_bytes=None, resultado_sens=None):
    """
    Genera un reporte Excel (.xlsx) profesional con los resultados del modelo AHP.

    Parámetros
    ----------
    pesos_criterios : array-like, pesos globales de los criterios.
    criterios : list[str].
    df_ranking : DataFrame con columnas "Puntaje Global" y "Ranking", indexado por alternativa.
    mejor_opcion : str, nombre de la mejor alternativa.
    modo : str, descripción del modo usado (para mostrar en el resumen).
    fig_bytes : bytes o None. PNG del gráfico (pesos + ranking) para incrustar, opcional.
    resultado_sens : dict o None. Resultado de ejecutar_analisis_sensibilidad(), opcional.

    Retorna: bytes del archivo .xlsx, listos para descargar (st.download_button).
    """
    wb = Workbook()

    # ---------------------------------------------------------------------------------
    # HOJA 1: RESUMEN EJECUTIVO
    # ---------------------------------------------------------------------------------
    ws = wb.active
    ws.title = "Resumen"

    ws["B2"] = "Sistema de Soporte a la Decisión Multicriterio (AHP)"
    ws["B2"].font = Font(name=_FUENTE, bold=True, size=16, color=_AZUL_OSCURO)

    ws["B4"] = "Modo de análisis:"
    ws["B4"].font = Font(name=_FUENTE, bold=True, size=11)
    ws["C4"] = modo
    ws["C4"].font = Font(name=_FUENTE, size=11)

    ws["B5"] = "Fecha del reporte:"
    ws["B5"].font = Font(name=_FUENTE, bold=True, size=11)
    ws["C5"] = datetime.now().strftime("%d/%m/%Y %H:%M")
    ws["C5"].font = Font(name=_FUENTE, size=11)

    ws["B7"] = "🏆 Mejor alternativa recomendada:"
    ws["B7"].font = Font(name=_FUENTE, bold=True, size=12, color=_AZUL_OSCURO)
    ws["C7"] = str(mejor_opcion)
    ws["C7"].font = Font(name=_FUENTE, bold=True, size=12, color=_VERDE)

    ws["B8"] = "Puntaje global obtenido:"
    ws["B8"].font = Font(name=_FUENTE, bold=True, size=11)
    ws["C8"] = float(df_ranking.iloc[0]["Puntaje Global"])
    ws["C8"].number_format = "0.0000"
    ws["C8"].font = Font(name=_FUENTE, size=11)

    texto_conclusion = (
        f"De acuerdo con el modelo AHP aplicado ({modo}), la alternativa mejor evaluada es "
        f"\"{mejor_opcion}\", con un puntaje global de {df_ranking.iloc[0]['Puntaje Global']:.4f} "
        f"sobre un total de {len(df_ranking)} alternativas comparadas. Este resultado surge de "
        f"ponderar el desempeño de cada alternativa en cada criterio según la importancia relativa "
        f"asignada a los {len(criterios)} criterios evaluados: {', '.join(criterios)}."
    )
    ws["B10"] = "Conclusión:"
    ws["B10"].font = Font(name=_FUENTE, bold=True, size=11)
    ws["B11"] = texto_conclusion
    ws["B11"].font = Font(name=_FUENTE, size=10, italic=True)
    ws["B11"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("B11:H16")

    if resultado_sens is not None:
        fila_sens = 18
        ws.cell(row=fila_sens, column=2, value="Análisis de sensibilidad:").font = Font(
            name=_FUENTE, bold=True, size=11
        )
        estado_txt = (
            "el ranking se mantuvo ESTABLE ante el cambio de pesos aplicado."
            if resultado_sens["ranking_estable"]
            else "el ranking CAMBIÓ ante el cambio de pesos aplicado."
        )
        ws.cell(row=fila_sens + 1, column=2,
                value=f"Se evaluó un escenario alternativo de pesos y {estado_txt}").font = Font(
            name=_FUENTE, size=10, italic=True
        )
        ws.cell(row=fila_sens + 1, column=2).alignment = Alignment(wrap_text=True)
        ws.merge_cells(start_row=fila_sens + 1, start_column=2, end_row=fila_sens + 2, end_column=8)

    for col in "ABCDEFGH":
        ws.column_dimensions[col].width = 16

    if fig_bytes is not None:
        img = XLImage(io.BytesIO(fig_bytes))
        img.width, img.height = 720, 300
        ws.add_image(img, "B22")

    # ---------------------------------------------------------------------------------
    # HOJA 2: RANKING FINAL DE ALTERNATIVAS
    # ---------------------------------------------------------------------------------
    ws2 = wb.create_sheet("Ranking Final")
    ws2["A1"] = "Ranking Final de Alternativas"
    ws2["A1"].font = Font(name=_FUENTE, bold=True, size=13, color=_AZUL_OSCURO)
    df_ranking_ordenado = df_ranking.sort_values("Ranking")
    _escribir_df(ws2, df_ranking_ordenado, fila_inicio=3, nombre_index="Alternativa",
                 resaltar_primera_fila=True)

    # ---------------------------------------------------------------------------------
    # HOJA 3: PESOS DE LOS CRITERIOS
    # ---------------------------------------------------------------------------------
    ws3 = wb.create_sheet("Pesos de Criterios")
    ws3["A1"] = "Vector de Pesos Globales de los Criterios"
    ws3["A1"].font = Font(name=_FUENTE, bold=True, size=13, color=_AZUL_OSCURO)
    df_pesos_crit = pd.DataFrame({"Peso (W)": pesos_criterios}, index=criterios)
    _escribir_df(ws3, df_pesos_crit, fila_inicio=3, nombre_index="Criterio")

    # ---------------------------------------------------------------------------------
    # HOJA 4: ANÁLISIS DE SENSIBILIDAD (solo si se ejecutó)
    # ---------------------------------------------------------------------------------
    if resultado_sens is not None:
        ws4 = wb.create_sheet("Sensibilidad")
        ws4["A1"] = "Análisis de Sensibilidad — Comparativa de Pesos"
        ws4["A1"].font = Font(name=_FUENTE, bold=True, size=13, color=_AZUL_OSCURO)
        fila_sig = _escribir_df(
            ws4, resultado_sens["df_comparativo_criterios"], fila_inicio=3, nombre_index="Criterio"
        )

        ws4.cell(row=fila_sig, column=1, value="Impacto en el Ranking de Alternativas").font = Font(
            name=_FUENTE, bold=True, size=13, color=_AZUL_OSCURO
        )
        fila_sig = _escribir_df(
            ws4, resultado_sens["df_resultados"], fila_inicio=fila_sig + 2, nombre_index="Alternativa"
        )

        estado_txt = (
            "El ranking se mantuvo ESTABLE ante el cambio de pesos."
            if resultado_sens["ranking_estable"]
            else "El ranking CAMBIÓ ante el cambio de pesos."
        )
        ws4.cell(row=fila_sig, column=1, value=f"Conclusión: {estado_txt}").font = Font(
            name=_FUENTE, bold=True, italic=True, size=11
        )

    # ---------------------------------------------------------------------------------
    # GUARDAR EN MEMORIA Y RETORNAR BYTES
    # ---------------------------------------------------------------------------------
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
